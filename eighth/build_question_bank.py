#!/usr/bin/env python3
"""Extract the decks' question bank and inject it into IT250_Study_Guide.md.

The six .pptx decks carry ~5,500 words of assessment that the study guide did
not: a two-question Quick Check after every session, an end-of-unit
multiple-choice recap, and end-of-unit short-answer / applied-case / discussion
sections. On a slide the correct option is marked inline with a tick, which is
useless for self-testing on paper - so this regroups each unit into a question
set followed by a separate answer key.

It is idempotent: re-running replaces the generated sections rather than
stacking another copy, so the decks stay the source of truth.

Usage
    cd eighth && /usr/bin/python3 build_question_bank.py
    /usr/bin/python3 build_question_bank.py --check   # report, write nothing
"""

import argparse
import pathlib
import re
import sys

from pptx import Presentation

ROOT = pathlib.Path(__file__).resolve().parent
GUIDE = ROOT / "IT250_Study_Guide.md"
DECKS = [ROOT / ("IT250_Unit%d.pptx" % u) for u in range(1, 7)]

TICK = "✅"
BEGIN = "<!-- question-bank:unit%d:begin -->"
END = "<!-- question-bank:unit%d:end -->"

OPT_RE = re.compile(r"(?:^|\s)([a-d])\)\s*")
QSTEM_RE = re.compile(r"^Q\d+\.\s*(.*)$")
RECALL_RE = re.compile(r"^\d+\.\s*(.*?)\s*(?:→|->)\s*%s?\s*(.*)$" % TICK)
SPLIT_NUMBERED = re.compile(r"\s{2,}(?=\d+\.\s)")


def slide_lines(slide):
    """Every line of text on a slide, tables included, page number dropped."""
    out = []
    for shape in slide.shapes:
        if shape.has_text_frame:
            out += shape.text_frame.text.split("\n")
        if shape.has_table:
            for row in shape.table.rows:
                out.append(" | ".join(c.text.strip().replace("\n", " ")
                                      for c in row.cells))
    return [l.rstrip() for l in out
            if l.strip() and not re.fullmatch(r"\d+\s*/\s*\d+", l.strip())]


def parse_options(line):
    """'a) x  b) [tick] y  c) z' -> ([x, y, z], index of the ticked one)."""
    parts = OPT_RE.split(line.strip())
    if len(parts) < 3:
        return None, None
    opts, answer = [], None
    # split() yields ['', 'a', text, 'b', text, ...]
    for letter, text in zip(parts[1::2], parts[2::2]):
        text = text.strip()
        if TICK in text:
            answer = len(opts)
            text = text.replace(TICK, "").strip()
        opts.append(text)
    return opts, answer


def parse_quick_check(lines):
    """A Quick Check slide -> its MCQs and its discussion prompt."""
    mcqs, discussion = [], None
    i = 0
    while i < len(lines):
        m = QSTEM_RE.match(lines[i].strip())
        if m:
            stem, opts, answer, why = m.group(1).strip(), None, None, None
            j = i + 1
            while j < len(lines) and not QSTEM_RE.match(lines[j].strip()):
                body = lines[j].strip()
                if opts is None and OPT_RE.search(body):
                    opts, answer = parse_options(body)
                elif body.startswith("Why:"):
                    why = body[4:].strip()
                elif body.startswith("Discussion:"):
                    discussion = body[len("Discussion:"):].strip()
                j += 1
            if opts and answer is not None:
                mcqs.append({"stem": stem, "options": opts,
                             "answer": answer, "why": why})
            i = j
            continue
        if lines[i].strip().startswith("Discussion:"):
            discussion = lines[i].strip()[len("Discussion:"):].strip()
        i += 1
    return mcqs, discussion


def parse_recall(lines):
    """Section A, which comes in two shapes across the decks: a 'stem -> answer'
    recall list, or a full a)-d) multiple choice. Both are kept, tagged."""
    out = []
    for line in lines:
        s = line.strip()
        if not re.match(r"^\d+\.\s", s):
            continue
        m = re.match(r"^\d+\.\s*(.*?):?\s{2,}(?=[a-d]\))(.*)$", s)
        if m and OPT_RE.search(m.group(2)):
            opts, answer = parse_options(m.group(2))
            if opts and answer is not None:
                out.append({"kind": "mcq", "stem": m.group(1).strip(),
                            "options": opts, "answer": answer, "why": None})
                continue
        m = RECALL_RE.match(s)
        if m and m.group(2).strip():
            out.append({"kind": "recall", "stem": m.group(1).strip(),
                        "answer": m.group(2).strip()})
    return out


def parse_sections_bcd(lines):
    """Sections B/C/D -> {'B': [q, ...], 'C': [...], 'D': [...]}."""
    out, current = {"B": [], "C": [], "D": []}, None
    for line in lines:
        s = line.strip()
        m = re.match(r"^Sections?\s+([B-D])\b", s)
        if m:
            current = m.group(1)
            rest = s[m.end():].lstrip(" —-–").strip()
            if not re.match(r"^\d+\.", rest):
                continue
            s = rest
        if current and re.match(r"^\d+\.\s", s):
            for part in SPLIT_NUMBERED.split(s):
                part = re.sub(r"^\d+\.\s*", "", part.strip())
                if part:
                    out[current].append(part)
    return out


def harvest(deck_path):
    """Pull every assessment item out of one unit's deck."""
    prs = Presentation(deck_path)
    mcqs, discussions, recall, bcd = [], [], [], {"B": [], "C": [], "D": []}
    for slide in prs.slides:
        lines = slide_lines(slide)
        blob = " ".join(lines).upper()
        if "CHECK FOR UNDERSTANDING" not in blob and "QUICK CHECK" not in blob:
            continue
        if "QUICK CHECK" in blob:
            qs, disc = parse_quick_check(lines)
            mcqs += qs
            if disc:
                discussions.append(disc)
        elif "SECTION A" in blob:
            recall += parse_recall(lines)
        else:
            got = parse_sections_bcd(lines)
            for k in bcd:
                bcd[k] += got[k]
    return mcqs, discussions, recall, bcd


def sentence_case(text):
    """The slides' discussion prompts start mid-sentence ('name a Nepali ...')
    because a label preceded them; on the page they stand alone."""
    return text[:1].upper() + text[1:] if text else text


def render(unit, title, mcqs, discussions, recall, bcd):
    """Build the markdown for one unit's test section."""
    L = ["## %d.T Test yourself — Unit %d" % (unit, unit), ""]
    L.append("Work through the questions first and keep the answer key covered. "
             "Parts A and B have fixed answers, given at the end of this section; "
             "Parts C to E are the open questions that carry the most marks in an "
             "exam, so practise writing them out in full.")
    L.append("")

    if mcqs:
        L += ["### Part A — Multiple choice", ""]
        for n, q in enumerate(mcqs, 1):
            L.append("**A%d.** %s" % (n, q["stem"]))
            L.append("")
            L.append("  " + "  ·  ".join(
                "*%s)* %s" % (chr(97 + i), o) for i, o in enumerate(q["options"])))
            L.append("")

    if recall:
        mixed = any(q["kind"] == "mcq" for q in recall)
        L += ["### Part B — Unit review", "",
              "The whole unit in one pass."
              if mixed else "Complete each line from memory.", ""]
        for n, q in enumerate(recall, 1):
            if q["kind"] == "mcq":
                L.append("**B%d.** %s" % (n, q["stem"]))
                L.append("")
                L.append("  " + "  ·  ".join(
                    "*%s)* %s" % (chr(97 + i), o)
                    for i, o in enumerate(q["options"])))
            else:
                L.append("**B%d.** %s → ?" % (n, q["stem"]))
            L.append("")

    labels = [("C", "Part C — Short answer", bcd["B"]),
              ("D", "Part D — Applied case", bcd["C"]),
              ("E", "Part E — Discussion", bcd["D"] + discussions)]
    for tag, heading, items in labels:
        if not items:
            continue
        L += ["### " + heading, ""]
        for n, q in enumerate(items, 1):
            L.append("**%s%d.** %s" % (tag, n, sentence_case(q)))
            L.append("")

    if mcqs or recall:
        L += ["### Answers — Unit %d" % unit, ""]
        if mcqs:
            L += ["| # | Answer | Why |", "|---|---|---|"]
            for n, q in enumerate(mcqs, 1):
                letter = chr(97 + q["answer"])
                L.append("| **A%d** | **%s)** %s | %s |"
                         % (n, letter, q["options"][q["answer"]],
                            sentence_case(q["why"]) if q["why"] else "—"))
            L.append("")
        if recall:
            L += ["| # | Answer |", "|---|---|"]
            for n, q in enumerate(recall, 1):
                ans = ("**%s)** %s" % (chr(97 + q["answer"]),
                                       q["options"][q["answer"]])
                       if q["kind"] == "mcq" else q["answer"])
                L.append("| **B%d** | %s |" % (n, ans))
            L.append("")
    return "\n".join(L).rstrip() + "\n"


def inject(text, unit, body):
    """Replace this unit's generated block, or add it at the unit's end."""
    begin, end = BEGIN % unit, END % unit
    block = "%s\n%s\n%s\n" % (begin, body, end)
    if begin in text:
        return re.sub(re.escape(begin) + r".*?" + re.escape(end) + r"\n?",
                      block, text, flags=re.S)
    # the unit ends at the rule that precedes the next "# Unit" / "# Course-Wide"
    nxt = re.search(r"\n---\n+(?=# (?:Unit %d|Course-Wide))" % (unit + 1), text)
    if not nxt:
        sys.exit("Could not find where Unit %d ends." % unit)
    return text[:nxt.start()] + "\n\n" + block + text[nxt.start():]


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--check", action="store_true",
                    help="report what would be extracted, write nothing")
    args = ap.parse_args()

    text = GUIDE.read_text(encoding="utf-8")
    titles = dict(re.findall(r"^# (Unit (\d)[^\n]*)$", text, re.M))
    total = 0
    for unit, deck in enumerate(DECKS, 1):
        if not deck.exists():
            sys.exit("Missing deck: %s" % deck)
        mcqs, discussions, recall, bcd = harvest(deck)
        n = len(mcqs) + len(recall) + sum(len(v) for v in bcd.values()) \
            + len(discussions)
        total += n
        print("  Unit %d: %2d MCQ · %2d recall · %d short · %d case · %d discussion"
              % (unit, len(mcqs), len(recall), len(bcd["B"]), len(bcd["C"]),
                 len(bcd["D"]) + len(discussions)))
        if not args.check:
            body = render(unit, titles.get(str(unit), ""), mcqs, discussions,
                          recall, bcd)
            text = inject(text, unit, body)

    if args.check:
        print("  %d questions found (nothing written)" % total)
        return
    GUIDE.write_text(text, encoding="utf-8")
    print("  %d questions -> %s" % (total, GUIDE.name))


if __name__ == "__main__":
    main()
